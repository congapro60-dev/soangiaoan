"""Read the saved PPCT workbook. Never modify it or import HKI overrides."""

import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries


SOURCE = Path('C:/Users/ADMIN/Downloads/PPCT THCS-THPT 26-27.xlsx')
OUT = Path(__file__).parent
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
NAMES = ['G_Lich năm học', 'G10', 'G11', 'G12']
KINDS = ['core', 'elective', 'review', 'assessment', 'assessment_feedback',
         'experiential', 'summer_assignment', 'orientation', 'unspecified']


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def read_source():
    result = {}
    with ZipFile(SOURCE) as archive:
        strings = [''.join(si.itertext()) for si in
                   ET.fromstring(archive.read('xl/sharedStrings.xml'))]
        rels = {e.attrib['Id']: e.attrib['Target'] for e in
                ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        for sheet in ET.fromstring(archive.read('xl/workbook.xml')).find('s:sheets', NS):
            name = sheet.attrib['name']
            if name not in NAMES:
                continue
            target = rels[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            path = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            root = ET.fromstring(archive.read(path))
            cells = {}
            for cell in root.findall('.//s:sheetData/s:row/s:c', NS):
                ref = cell.attrib['r']
                col = re.sub(r'\d+', '', ref)
                if name != NAMES[0] and col not in ['A', 'B', 'C', 'D', 'E', 'H']:
                    continue
                value = cell.find('s:v', NS)
                formula = cell.find('s:f', NS)
                text = value.text if value is not None else None
                kind = cell.attrib.get('t', 'n')
                if kind == 's' and text is not None:
                    cached = strings[int(text)]
                elif kind == 'inlineStr':
                    cached = ''.join(cell.find('s:is', NS).itertext())
                elif kind == 'b' and text is not None:
                    cached = text == '1'
                elif kind == 'n' and text is not None:
                    cached = float(text)
                    if cached.is_integer():
                        cached = int(cached)
                else:
                    cached = text
                raw = '=' + (formula.text or '') if formula is not None else cached
                if raw is not None or cached is not None:
                    cells[ref] = {'raw': raw, 'cached': cached,
                                  'xml_type': kind, 'style_id': cell.attrib.get('s')}
            result[name] = {
                'xml_path': path, 'cells': cells,
                'merges': sorted(e.attrib['ref'] for e in root.findall('s:mergeCells/s:mergeCell', NS)),
            }
    return result


def resolver(name, sheet):
    merged = {}
    for area in sheet['merges']:
        left, top, right, bottom = range_boundaries(area)
        anchor = f'{get_column_letter(left)}{top}'
        for row in range(top, bottom + 1):
            for col in range(left, right + 1):
                merged[f'{get_column_letter(col)}{row}'] = (anchor, area)

    def get(coordinate):
        anchor, area = merged.get(coordinate, (coordinate, coordinate))
        item = sheet['cells'].get(anchor, {'raw': None, 'cached': None})
        return {'cell': coordinate, 'anchor': anchor, 'range': area,
                'ref': f'{name}!{area}', **item}

    return get


def classify(content, subject):
    text = content.casefold()
    if 'tự chọn' in text or (not text and 'tự chọn' in subject.casefold()):
        return 'elective'
    if 'kiểm tra' in text and ('chữa' in text or 'trả' in text):
        return 'assessment_feedback'
    if text.startswith('kiểm tra'):
        return 'assessment'
    if text.startswith(('ôn tập', 'ôn thi', 'tổng ôn', 'luyện đề')):
        return 'review'
    if text.startswith(('hoạt động trải nghiệm', 'hoạt động dự án', 'báo cáo dự án')):
        return 'experiential'
    if text.startswith('giao nhiệm vụ bài tập hè'):
        return 'summer_assignment'
    if text.startswith('giới thiệu'):
        return 'orientation'
    return 'core' if text else 'unspecified'


def extract_grade(grade, sheet):
    name = f'G{grade}'
    get = resolver(name, sheet)
    all_lessons, unnumbered, missing_week, contexts = [], [], [], {}
    week_anchors = sorted((int(re.sub(r'\D', '', ref)), int(c['cached']))
                          for ref, c in sheet['cells'].items()
                          if re.fullmatch(r'B\d+', ref) and numeric(c['cached']) and 1 <= c['cached'] <= 53)
    last_week = max(w for _, w in week_anchors)
    last_row = range_boundaries(get(f'B{week_anchors[-1][0]}')['range'])[3]
    for row in range(week_anchors[0][0], last_row + 1):
        source = {key: get(f'{col}{row}') for key, col in
                  [('week', 'B'), ('subject', 'C'), ('period', 'D'), ('content', 'E'), ('notes', 'H')]}
        week, period = source['week']['cached'], source['period']['cached']
        content = str(source['content']['cached'] or '').strip()
        subject = str(source['subject']['cached'] or '').strip()
        if not numeric(week):
            if numeric(period) or content:
                missing_week.append({'row': row, 'source': source})
            continue
        if not numeric(period):
            # Content anchors still matter when the D merge has no number (G11 projects).
            if week >= 19 and content and source['content']['anchor'] == f'E{row}':
                unnumbered.append({'row': row, 'week': int(week), 'kind': classify(content, subject),
                                   'content': content, 'source': source, 'period_count': None})
            continue
        if source['period']['anchor'] != f'D{row}':
            continue
        assert int(period) == period and period > 0
        kind = classify(content, subject)
        lines = [line.strip().strip('"') for line in content.splitlines() if line.strip().strip('"')]
        first = lines[0] if lines else None
        topic, basis = first, {'mode': 'explicit', 'ref': source['content']['ref']}
        if kind == 'elective':
            topic = 'Tự chọn'
            basis = {'mode': 'explicit' if content else 'subject_label',
                     'ref': source['content' if content else 'subject']['ref']}
        elif first and re.match(r'^Tiết\s*\d+', first, re.I):
            previous = contexts.get(subject)
            topic = previous['topic'] if previous else subject.splitlines()[0] if subject else first
            basis = {'mode': 'preceding_topic_context' if previous else 'subject_context',
                     'ref': previous['ref'] if previous else source['subject']['ref']}
        if kind == 'core' and first and not re.match(r'^Tiết\s*\d+', first, re.I):
            contexts[subject] = {'topic': topic, 'ref': source['content']['ref']}
        all_lessons.append({
            'id': f'g{grade}-discover-r{row:03}', 'grade': grade, 'track': 'Discover',
            'week': int(week), 'period': int(period), 'row': row,
            'subject': subject.splitlines()[0] if subject else None,
            'topic': topic, 'detail': '\n'.join(lines[1:] if first == topic else lines),
            'kind': kind, 'effective_kind': kind,
            'activity_type': 'cumulative_practice_game' if kind == 'elective' else kind,
            'practice_eligible': kind == 'core', 'topic_basis': basis,
            'source': {'sheet': name, **source},
        })
    lessons = [p for p in all_lessons if p['week'] >= 19]
    prior = [p for p in all_lessons if p['week'] < 19 and p['practice_eligible']]
    core = [p for p in lessons if p['practice_eligible']]
    for seq, lesson in enumerate(lessons, 1):
        lesson['sequence'] = seq
    scopes, weeks = [], []
    final_exams = [p for p in lessons if p['kind'] == 'assessment' and 'cuối' in p['topic']]
    final_week = max(p['week'] for p in final_exams)
    for week in range(19, last_week + 1):
        weekly = [p for p in lessons if p['week'] == week]
        past = prior + [p for p in core if p['week'] < week]
        current = [p for p in weekly if p['practice_eligible']]
        scope = {'id': f'g{grade}-discover-w{week:02}', 'week': week,
                 'prior_week_lesson_ids': [p['id'] for p in past],
                 'prior_week_topics': list(dict.fromkeys(p['topic'] for p in past)),
                 'same_week_candidate_lesson_ids': [p['id'] for p in current],
                 'same_week_candidate_topics': list(dict.fromkeys(p['topic'] for p in current)),
                 'same_week_order_verified': False}
        scopes.append(scope)
        for lesson in weekly:
            lesson['phase'] = 'through_final_exam_week' if week <= final_week else 'after_final_exam_week'
            if lesson['kind'] == 'elective':
                lesson['elective_scope'] = {
                    'practice_scope_id': scope['id'], 'mode': 'cumulative_practice',
                    'same_week_earlier_lesson_ids': [p['id'] for p in current if p['row'] < lesson['row']],
                    'same_week_later_lesson_ids': [p['id'] for p in current if p['row'] > lesson['row']],
                    'same_week_order_verified': False,
                    'rule': 'Default: prior-week source details only. Same-week earlier rows require teacher confirmation of actual teaching. Later rows and future weeks are excluded.',
                }
        anchors = [get(f'B{r}') for r, w in week_anchors if w == week]
        weeks.append({'week': week, 'count': len(weekly), 'periods': [p['period'] for p in weekly],
                      'kinds': dict(Counter(p['kind'] for p in weekly)), 'source': anchors,
                      'unnumbered_content_count': len([p for p in unnumbered if p['week'] == week])})
    numbers = [p['period'] for p in lessons]
    no_period_electives = []
    for ref, cell in sheet['cells'].items():
        if not re.fullmatch(r'C\d+', ref) or 'tự chọn' not in str(cell['cached']).casefold():
            continue
        row = int(ref[1:])
        week = get(f'B{row}')['cached']
        if not numeric(week) or week < 19:
            continue
        _, top, _, bottom = range_boundaries(get(ref)['range'])
        if not any(p['row'] in range(top, bottom + 1) for p in lessons):
            no_period_electives.append({'week': int(week), 'source': get(ref),
                                       'period_cells': [get(f'D{r}') for r in range(top, bottom + 1)]})
    return {'grade': grade, 'summary': {
        'actual_period_rows': len(lessons), 'first_week': 19, 'last_week': last_week,
        'source_first_period': min(numbers), 'source_last_period': max(numbers),
        'kinds': {kind: sum(p['kind'] == kind for p in lessons) for kind in KINDS},
        'duplicate_period_numbers': {str(n): [p['id'] for p in lessons if p['period'] == n]
                                     for n, count in Counter(numbers).items() if count > 1},
        'missing_period_numbers': sorted(set(range(min(numbers), max(numbers) + 1)) - set(numbers)),
        'final_exam_week': final_week,
        'through_final_exam_week_count': sum(p['week'] <= final_week for p in lessons),
        'after_final_exam_week_count': sum(p['week'] > final_week for p in lessons),
        'week19_prior_semester_feedback_count': sum(p['week'] == 19 and p['kind'] == 'assessment_feedback' for p in lessons),
        'unspecified_periods': [p['source']['period']['ref'] for p in lessons if p['kind'] == 'unspecified'],
        'content_without_period_count': len(unnumbered), 'numbered_or_content_without_week': missing_week,
        'elective_sections_without_numbered_periods': no_period_electives,
        'zero_numbered_period_weeks': [w['week'] for w in weeks if not w['count']],
    }, 'weeks': weeks, 'lessons': lessons, 'practice_scopes': scopes,
       'unnumbered_activity_rows': unnumbered,
       'hki_prerequisites': {'mode': 'read_only_source_reextraction', 'owner_overrides_imported': False,
                             'lessons': prior, 'topics': list(dict.fromkeys(p['topic'] for p in prior))},
       'merged_ranges': sheet['merges']}


def split_cutoff(grade):
    lessons = grade['lessons']
    final = [p for p in lessons if p['kind'] == 'assessment' and 'cuối HK2' in (p['topic'] or '')]
    feedback = [p for p in lessons if p['kind'] == 'assessment_feedback' and 'cuối HK2' in (p['topic'] or '')]
    last = max(feedback or final, key=lambda p: p['row'])
    primary = [p for p in lessons if p['row'] <= last['row']]
    after = [p for p in lessons if p['row'] > last['row']]
    unnumbered = grade['unnumbered_activity_rows']
    grade['cutoff'] = {'row': last['row'], 'week': last['week'], 'period': last['period'],
                       'ref': last['source']['content']['ref'],
                       'basis': 'last_explicit_final_exam_feedback' if feedback else 'last_explicit_final_exam_no_feedback_in_source',
                       'final_exam_feedback_present': bool(feedback),
                       'same_week_rows_after_cutoff_are_post_hkii': True}
    grade['full_remainder_summary'] = grade['summary']
    grade['post_hkii'] = {'author_in_current_task': False, 'lessons': after,
                         'unnumbered_activity_rows': [p for p in unnumbered if p['row'] > last['row']],
                         'practice_scopes': [s for s in grade['practice_scopes']
                                             if s['week'] in {p['week'] for p in after}],
                         'summary': {'actual_period_rows': len(after),
                                     'kinds': dict(Counter(p['kind'] for p in after))}}
    grade['lessons'] = primary
    grade['unnumbered_activity_rows'] = [p for p in unnumbered if p['row'] <= last['row']]
    grade['practice_scopes'] = [s for s in grade['practice_scopes'] if s['week'] <= last['week']]
    for scope in grade['practice_scopes']:
        scope['same_week_candidate_lesson_ids'] = [i for i in scope['same_week_candidate_lesson_ids']
                                                  if i in {p['id'] for p in primary}]
    for week in grade['weeks']:
        rows = [p for p in primary if p['week'] == week['week']]
        week.update(count=len(rows), periods=[p['period'] for p in rows],
                    kinds=dict(Counter(p['kind'] for p in rows)),
                    unnumbered_content_count=sum(p['week'] == week['week'] for p in grade['unnumbered_activity_rows']))
    grade['weeks'] = [w for w in grade['weeks'] if w['week'] <= last['week']]
    numbers = [p['period'] for p in primary]
    grade['summary'] = {
        'actual_period_rows': len(primary), 'first_week': 19, 'last_week': last['week'],
        'source_first_period': min(numbers), 'source_last_period': max(numbers),
        'kinds': {kind: sum(p['kind'] == kind for p in primary) for kind in KINDS},
        'missing_period_numbers': sorted(set(range(min(numbers), max(numbers) + 1)) - set(numbers)),
        'duplicate_period_numbers': {str(n): [p['id'] for p in primary if p['period'] == n]
                                     for n, count in Counter(numbers).items() if count > 1},
        'unspecified_periods': [p['source']['period']['ref'] for p in primary if p['kind'] == 'unspecified'],
        'content_without_period_count': len(grade['unnumbered_activity_rows']),
        'zero_numbered_period_weeks': [w['week'] for w in grade['weeks'] if not w['count']],
        'elective_sections_without_numbered_periods': [s for s in grade['full_remainder_summary']['elective_sections_without_numbered_periods']
                                                     if int(re.sub(r'\D', '', s['source']['cell'])) <= last['row']],
    }
    return grade


def validate(result, raw):
    cached = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    formulas = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False)
    checked, formula_count = {}, 0
    for grade in result['grades']:
        name = f"G{grade['grade']}"
        lessons = grade['lessons'] + grade['post_hkii']['lessons']
        by_id = {p['id']: p for p in grade['hki_prerequisites']['lessons'] + lessons}
        assert len(by_id) == len(grade['hki_prerequisites']['lessons']) + len(lessons)
        assert not grade['full_remainder_summary']['numbered_or_content_without_week']
        get = resolver(name, raw[name])
        low = min(p['row'] for p in lessons)
        high = max(int(re.sub(r'\D', '', ref)) for ref in raw[name]['cells'])
        expected = set()
        for row_number, row in enumerate(cached[name].iter_rows(min_row=low, max_row=high, max_col=8), low):
            for index in [1, 2, 3, 4, 7]:
                cell = row[index]
                coordinate = f'{get_column_letter(index + 1)}{row_number}'
                direct = raw[name]['cells'].get(coordinate, {})
                assert cell.value == direct.get('cached'), (name, coordinate, cell.value, direct)
            if numeric(row[3].value):
                expected.add((row[3].row, int(row[3].value)))
        assert {(p['row'], p['period']) for p in lessons} == expected
        checked[name] = len(expected)
        for row in formulas[name].iter_rows(min_row=low, max_row=high, min_col=4, max_col=4):
            cell = row[0]
            if cell.data_type != 'f':
                continue
            match = re.fullmatch(r'=ROW\(\)-(\d+)', cell.value)
            assert match and cell.row - int(match[1]) == get(cell.coordinate)['cached'], (name, cell.coordinate)
            formula_count += 1
        scopes = grade['practice_scopes'] + grade['post_hkii']['practice_scopes']
        for scope in scopes:
            assert all(by_id[i]['week'] < scope['week'] and by_id[i]['practice_eligible']
                       for i in scope['prior_week_lesson_ids'])
            assert all(by_id[i]['week'] == scope['week'] and by_id[i]['practice_eligible']
                       for i in scope['same_week_candidate_lesson_ids'])
        for lesson in lessons:
            assert 'ownerOverride' not in lesson
            assert lesson['kind'] == lesson['effective_kind']
            if lesson['kind'] == 'elective':
                assert lesson['elective_scope']['mode'] == 'cumulative_practice'
                assert lesson['elective_scope']['practice_scope_id'] in {s['id'] for s in scopes}
                assert all(by_id[i]['row'] < lesson['row'] for i in lesson['elective_scope']['same_week_earlier_lesson_ids'])
    cached.close()
    formulas.close()
    assert [g['full_remainder_summary']['actual_period_rows'] for g in result['grades']] == [158, 138, 233]
    assert [g['summary']['actual_period_rows'] for g in result['grades']] == [130, 130, 156]
    assert next(p for p in result['grades'][0]['post_hkii']['lessons'] if p['row'] == 377)['source']['content']['anchor'] == 'E376'
    assert [len(g['post_hkii']['unnumbered_activity_rows']) for g in result['grades']] == [0, 15, 0]
    assert not any(g['summary']['unspecified_periods'] for g in result['grades'])
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == result['source']['sha256']
    result['verification'] = {'source_sha256_unchanged': True, 'independent_openpyxl_period_rows': checked,
                              'primary_period_rows': [130, 130, 156], 'post_hkii_period_rows': [28, 8, 77],
                              'formula_rows_verified': formula_count, 'all_numbered_rows_exported_once': True,
                              'merged_content_E376_E377_retains_two_numbered_periods': True,
                              'no_owner_overrides_imported': True, 'no_future_week_in_practice_scopes': True}


def main():
    sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    raw = read_source()
    result = {
        'schema_version': 1,
        'source': {'path': str(SOURCE), 'sha256': sha, 'modified': False},
        'scope': {'semester': 'HKII', 'grades': [10, 11, 12], 'track': 'Discover', 'columns': 'D:E',
                  'first_week': 19, 'end': 'last explicit final-exam feedback; if absent, last final-exam row',
                  'after_cutoff': 'preserved separately in post_hkii; not authored in this task',
                  'owner_scope_instruction': 'Chỉ đến kiểm tra và trả bài cuối HKII; phần sau để riêng',
                  'excluded_tracks': ['Explore', 'DGS'],
                  'objectives_policy': 'PPCT identifies schedule and period details, not independently verified learning objectives.'},
        'boundary': {
            'status': 'inferred_content_boundary', 'official_semester_dates_confirmed': False,
            'basis': 'Week 18 ends HK1 examinations; week 19 contains HK1 feedback then new content. G12 changes to 10 numbered Discover periods in week19. Main scope ends at last explicit final-exam feedback, or last final-exam row when no feedback exists.',
            'evidence_refs': ["'G_Lich năm học'!B106:C106", 'G10!E167:E176',
                              'G11!E171:E180', 'G12!D18', 'G12!B183:E194',
                              'G10!E341', 'G11!E346:E347', 'G12!E402:E403'],
            'end_policy': 'G10 through week38 row341; G11 through week36 row346; G12 through week35 row403. G12 has no final-exam feedback. Its three elective rows404:406 are separate post_hkii despite sharing week35. No official dates inferred.',
        },
        'method': {'period': 'One numeric D anchor is one period; preserve original period number and source-row ID.',
                   'merged_cells': 'Resolve only within actual merges. Distinct numbered D anchors sharing one E merge remain distinct periods.',
                   'unnumbered_content': 'Keep separate source content anchors, including D merge children; never invent period counts.',
                   'electives': 'Cumulative practice game. Default prior-week core details, including source-extracted HKI. Same-week earlier content requires teacher confirmation.',
                   'review': 'On tap, On thi, Tong on and Luyen de are review, not formal assessment or new topic coverage.',
                   'source_dates': 'Date-styled XML numbers retained with style_id; no 2026-2027 calendar dates inferred.'},
        'ambiguities': [
            {'code': 'CALENDAR_MIXED_YEARS_AND_EXAM_WEEKS',
             'refs': ["'G_Lich năm học'!H5", "'G_Lich năm học'!B6", "'G_Lich năm học'!B184:C184", "'G_Lich năm học'!B220:C220", 'G11!B280:E283', 'G12!B323:E326'],
             'action': 'Use each grade sheet week, never calendar dates or a common exam week. Calendar says weeks31/37; G11/G12 examinations are weeks29/35.'},
            {'code': 'G10_PERIOD_GAPS', 'refs': ['G10!D320:D321', 'G10!D341:D347'],
             'action': 'Keep gaps256:258 in primary; gaps277:281 occur between primary and post_hkii. Do not create missing lessons.'},
            {'code': 'G12_PERIOD_GAP', 'refs': ['G12!D194:D205', 'G12!B195:H204'],
             'action': 'Keep gap148:157. Winter-break block has no numeric week or Discover periods; do not infer ten missing lessons.'},
            {'code': 'EMPTY_ELECTIVE_SECTIONS',
             'refs': ['G10!C187:E188', 'G10!C222:E223', 'G10!C318:E319', 'G11!C225:E226', 'G12!C214:E216', 'G12!C250:E252', 'G12!C380:E382'],
             'action': 'Subject says elective but D:E are blank. Preserve as sections without numbered periods, not extra games.'},
            {'code': 'G10_EMPTY_WEEK21', 'refs': ['G10!B189:E196'],
             'action': 'Week21 has zero numbered Discover periods and no explicit holiday label in this block. Do not invent content.'},
            {'code': 'G11_WEEK32_NINE_PERIODS', 'refs': ['G11!B307:E316', 'G11!D9'],
             'action': 'Retain nine numeric periods232:240 despite advertised eight per week.'},
            {'code': 'G11_REPEATED_DISTANCE_DETAIL', 'refs': ['G11!E304', 'G11!E309'],
             'action': 'Both source rows say lesson5 practice on distances. Keep distinct periods229 and234; do not deduplicate or invent lesson6.'},
            {'code': 'G11_GEOMETRY_REVIEW_CHAPTER_LABEL', 'refs': ['G11!C320:E323', 'G11!E246:E247'],
             'action': 'Geometry section week33 says chapterVII, also used in probability review week26. Do not infer exact review objectives from chapter number alone.'},
            {'code': 'G12_LINE_TOPIC_STARTS_AT_LESSON2', 'refs': ['G12!E233:E235', 'G12!E222:E225'],
             'action': 'No Discover row explicitly introduces lesson1 of spatial line equations. First is canonical equation lesson2. Do not silently add missing lesson1 or mark its detail already taught.'},
            {'code': 'G10_CONIC_SCOPE_ONLY_ELLIPSE_EXPLICIT', 'refs': ['G10!E239', 'G10!E246'],
             'action': 'Broad title Three conics has explicit ellipse lesson1 and practice lesson2 only. Do not treat hyperbola/parabola as learned merely from title.'},
            {'code': 'G12_FINAL_FEEDBACK_ABSENT', 'refs': ['G12!E402:E403', 'G12!C404:E410'],
             'action': 'Cut primary at final exam row403. No final feedback in source; preserve electives404:406 and all later rows in post_hkii.'},
            {'code': 'G11_UNNUMBERED_PROJECT_AND_CT12', 'refs': ['G11!D362:E370', 'G11!D371:E379', 'G11!D380:E382'],
             'action': 'Preserve15 distinct E content anchors. D362:D370 is blank; D371:D379 and D380:D382 say CT LỚP12. No period number or mathematical project scope inferred.'},
            {'code': 'G10_SUMMER_CONTENT_SHARED_MERGE', 'refs': ['G10!D376:E377'],
             'action': 'E376:E377 is one content merge shared by two numeric D anchors308/309; retain both post_hkii periods. This is not missing content.'},
            {'code': 'SAME_WEEK_ORDER_NOT_VERIFIED', 'refs': ['G10!B171:E341', 'G11!B175:E347', 'G12!B183:E403'],
             'action': 'Same-week source order is a candidate only. Elective games require teacher confirmation before using earlier same-week content.'},
            {'code': 'HKI_G10_TOPIC_DETAIL_CONFLICT', 'refs': ['G10!E30'],
             'action': 'Read-only prerequisite retains source conflict between trigonometric-value title/English and Vietnamese cosine-rule detail. Use exact Vietnamese detail with source review, never whole-topic expansion.'},
        ],
        'grades': [split_cutoff(extract_grade(grade, raw[f'G{grade}'])) for grade in [10, 11, 12]],
        'source_evidence': raw,
    }
    validate(result, raw)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / 'ppct-hkii.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({g['grade']: {k: v for k, v in g['summary'].items()
                                 if k not in ['content_without_period', 'elective_sections_without_numbered_periods']}
                      for g in result['grades']}, ensure_ascii=False, indent=2))
    print(json.dumps(result['verification'], ensure_ascii=False))


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
