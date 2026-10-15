import hashlib
import json
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from openpyxl.utils.cell import get_column_letter, range_boundaries


SOURCE = Path('C:/Users/ADMIN/Downloads/PPCT THCS-THPT 26-27.xlsx')
OUT = Path(__file__).parent


def inspect_workbook():
    before = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(SOURCE, data_only=False)
    cached = openpyxl.load_workbook(SOURCE, data_only=True)
    result = {'path': str(SOURCE), 'sha256': before, 'sheets': {}}
    for name in ['G_Lich năm học', 'G10', 'G11', 'G12']:
        sheet = workbook[name]
        result['sheets'][name] = {
            'merges': sorted(str(r) for r in sheet.merged_cells.ranges),
            'rows': [
                {'row': row[0].row, 'cells': {
                    c.coordinate: {'raw': c.value, 'cached': cached[name][c.coordinate].value}
                    for c in row if c.value is not None
                }}
                for row in sheet if any(c.value is not None for c in row)
            ],
        }
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == before, 'Source workbook changed'
    path = OUT / 'sources' / 'ppct-raw.json'
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    print(json.dumps({'sha256': before, 'rows': {k: len(v['rows']) for k, v in result['sheets'].items()}}, ensure_ascii=False))


def normalize():
    raw = json.loads((OUT / 'sources' / 'ppct-raw.json').read_text(encoding='utf-8'))
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == raw['sha256']
    result = {
        'schema_version': 1,
        'source': {'path': str(SOURCE), 'sha256': raw['sha256'], 'modified': False},
        'scope': {'grades': [10, 11, 12], 'track': 'Discover', 'columns': 'D:E',
                  'excluded_tracks': ['Explore', 'DGS'],
                  'objectives_policy': 'PPCT cung cấp tiến độ và nội dung tiết; không suy diễn YCCĐ/LO từ tên bài.'},
        'cutoff': {
            'last_week': 18,
            'basis': 'Mốc phạm vi nội dung: tuần kiểm tra cuối HK1; chưa có ngày kết thúc học kỳ chính thức trong nguồn.',
            'status': 'inferred_content_boundary',
            'evidence': [
                {'ref': "'G_Lich năm học'!B106:C106", 'value': 'Tuần 18; - Kiểm tra cuối HK1: 11/12'},
                {'ref': 'G10!E167:E168', 'value': 'Kiểm tra cuối HK1; Kiểm tra cuối HK1 (tiếp)'},
                {'ref': 'G11!E171:E172', 'value': 'Kiểm tra cuối HK1; Kiểm tra cuối HK1 (tiếp)'},
                {'ref': 'G12!E179:E180', 'value': 'Kiểm tra cuối HK1; Kiểm tra cuối HK1 (tiếp)'},
                {'ref': 'G12!D18', 'value': 'Discover: Kỳ 1: 8 tiết 1 tuần; Kỳ 2: 10 tiết 1 tuần.'},
                {'ref': 'G12!B183:D194', 'value': 'Tuần 19 có 10 tiết Discover, bắt đầu nội dung mới sau trả/chữa bài.'},
            ],
            'week19_feedback': 'Giữ riêng ở post_cutoff_feedback; không đếm vào tuần 1-18.',
            'dates': 'Không suy ngày 2026-2027: lịch chung còn ngày năm 2025-2026.'
        },
        'method': {
            'period_row': 'Một ô D có số tiết thực hoặc công thức trả số là một tiết. Không nhân bản các dòng con của ô D gộp.',
            'merged_cells': 'Giải tham chiếu B/C/E/H trong đúng vùng gộp; giữ anchor, range, raw, cached. Không fill-down qua ô trống ngoài vùng gộp.',
            'periods': 'Giữ nguyên số tiết nguồn; sequence là thứ tự dòng thật, không sửa lỗi đánh số của nguồn.',
            'electives': 'Ôn tập tích lũy các tiết core đã học; bài cùng tuần chỉ dùng sau khi GV xác nhận đã dạy. Không lấy tên học phần làm bằng chứng đã học.',
            'owner_overrides': 'Giữ kind và ô nguồn. effective_kind áp dụng xác nhận riêng của người dùng; ownerOverride lưu căn cứ.',
        },
        'ambiguities': [
            {'code': 'CALENDAR_YEAR_MISMATCH', 'refs': ["'G_Lich năm học'!H5", "'G_Lich năm học'!B6", 'G10!B1', 'G11!B7', 'G12!B7'],
             'note': 'Tiêu đề khối ghi 2026-2027 nhưng lịch chung còn 2025-2026. Không gán ngày cụ thể cho tuần.'},
            {'code': 'G10_PERIOD_GAPS', 'refs': ['G10!D28:D30', 'G10!D85:D86'],
             'note': 'Số tiết nhảy 15 sang 17, 66 sang 68; không có tiết nguồn 16 hoặc 67. Giữ ID theo ô.'},
            {'code': 'G10_WEEK3_LOAD', 'refs': ['G10!B26:E35', 'G10!D3'],
             'note': 'Tuần 3 có 9 tiết thật (5 nội dung, 4 tự chọn) dù Discover ghi 8 tiết/tuần; không tự xóa tiết.'},
            {'code': 'G12_BLANK_CONTENT', 'refs': ['G12!D163:E164', 'G12!D172:E173'],
             'note': 'Tiết 121/129: E trống, C gộp ghi Ôn thi HK I; kind nguồn giữ unspecified. Người dùng xác nhận ngày 2026-10-15 triển khai thành game ôn tập tích lũy; ownerOverride/effective_kind=elective.',
             'resolution': 'owner_confirmed_cumulative_practice'},
            {'code': 'G10_PROJECT_CONTENT_MISSING', 'refs': ['G10!E117:E118'],
             'note': 'Hai tiết Báo cáo dự án Toán 10 không nêu dự án/nội dung toán. Giữ core, activity_type=project_presentation, practice_eligible=false.'},
            {'code': 'G10_BILINGUAL_TOPIC_CONFLICT', 'refs': ['G10!E30'],
             'note': 'Tiêu đề ghi giá trị lượng giác, chi tiết Việt ghi Định lý cosin, tiếng Anh ghi luyện giá trị lượng giác. Giữ nguyên nguồn; dùng chi tiết Việt để đối soát nguồn bài, không suy LO từ tiêu đề.'},
            {'code': 'SAME_WEEK_ORDER_UNKNOWN', 'refs': ['G10!B7:E170', 'G11!B13:E174', 'G12!B19:E182'],
             'note': 'PPCT có thứ tự dòng/số tiết nhưng không có thời khóa biểu thực tế. Bài cùng tuần chỉ là ứng viên sau xác nhận đã dạy.'},
        ],
        'grades': [],
    }
    for grade in [10, 11, 12]:
        name = f'G{grade}'
        sheet = raw['sheets'][name]
        cells = {k: v for row in sheet['rows'] for k, v in row['cells'].items()}
        merged = {}
        for area in sheet['merges']:
            left, top, right, bottom = range_boundaries(area)
            anchor = f'{get_column_letter(left)}{top}'
            for row in range(top, bottom + 1):
                for col in range(left, right + 1):
                    merged[f'{get_column_letter(col)}{row}'] = (anchor, area)

        def source_cell(coordinate):
            anchor, area = merged.get(coordinate, (coordinate, coordinate))
            cell = cells.get(anchor, {'raw': None, 'cached': None})
            return {'cell': coordinate, 'anchor': anchor, 'range': area,
                    'ref': f'{name}!{area}', **cell}

        lessons = []
        feedback = []
        no_period = []
        topic_context = {}
        for raw_row in sheet['rows']:
            row = raw_row['row']
            week_source = source_cell(f'B{row}')
            week = week_source['cached']
            if not isinstance(week, (int, float)) or not 1 <= week <= 19:
                continue
            period_source = source_cell(f'D{row}')
            if period_source['anchor'] != f'D{row}':
                continue
            period = period_source['cached']
            content_source = source_cell(f'E{row}')
            subject_source = source_cell(f'C{row}')
            content = (content_source['cached'] or '').strip()
            subject = (subject_source['cached'] or '').strip()
            if not isinstance(period, (int, float)):
                if content and week <= 18:
                    no_period.append({'row': row, 'week': week, 'source': content_source})
                continue
            formula = period_source['raw']
            if isinstance(formula, str) and formula.startswith('='):
                match = re.fullmatch(r'=ROW\(\)-(\d+)', formula)
                assert match and row - int(match[1]) == period, (name, row, formula, period)
            text = content.casefold()
            if 'tự chọn' in text or (not content and 'tự chọn' in subject.casefold()):
                kind = 'elective'
            elif 'kiểm tra' in text and ('chữa' in text or 'trả' in text):
                kind = 'assessment_feedback'
            elif text.startswith('kiểm tra'):
                kind = 'assessment'
            elif text.startswith('ôn tập') or text.startswith('ôn thi'):
                kind = 'review'
            elif text.startswith('giới thiệu'):
                kind = 'orientation'
            elif not content:
                kind = 'unspecified'
            else:
                kind = 'core'
            lines = [line.strip().strip('"') for line in content.splitlines() if line.strip().strip('"')]
            first = lines[0] if lines else None
            topic = first
            topic_basis = {'mode': 'explicit', 'ref': content_source['ref']}
            if kind == 'elective':
                topic = 'Tự chọn'
                topic_basis = {'mode': 'subject_label' if not content else 'explicit',
                               'ref': subject_source['ref'] if not content else content_source['ref']}
            elif first and re.match(r'^Tiết\s*\d+', first, re.I):
                previous = topic_context.get(subject)
                topic = previous['topic'] if previous else subject.splitlines()[0] if subject else first
                topic_basis = {'mode': 'preceding_topic_context' if previous else 'subject_context',
                               'ref': previous['ref'] if previous else subject_source['ref']}
            if kind == 'core' and first and not re.match(r'^Tiết\s*\d+', first, re.I):
                topic_context[subject] = {'topic': topic, 'ref': content_source['ref']}
            lesson = {
                'id': f'g{grade}-discover-r{row:03}', 'grade': grade, 'track': 'Discover',
                'week': int(week), 'period': int(period), 'sequence': len(lessons) + 1,
                'row': row, 'subject': subject.splitlines()[0] if subject else None,
                'topic': topic, 'detail': '\n'.join(lines[1:] if first == topic else lines),
                'kind': kind, 'topic_basis': topic_basis,
                'practice_eligible': kind == 'core' and not text.startswith('báo cáo dự án'),
                'activity_type': 'project_presentation' if text.startswith('báo cáo dự án') else kind,
                'source': {'sheet': name, 'week': week_source, 'period': period_source,
                           'subject': subject_source, 'content': content_source, 'notes': source_cell(f'H{row}')},
            }
            lesson['effective_kind'] = kind
            if grade == 12 and row in [163, 172]:
                assert kind == 'unspecified' and not content
                lesson['ownerOverride'] = {
                    'kind': 'elective', 'activity_type': 'cumulative_practice_game',
                    'reason': 'ownerconfirmed2026-10-15', 'confirmed_by': 'owner',
                    'confirmed_date': '2026-10-15',
                    'instruction': 'Tiết 121 tuần 16 và tiết 129 tuần 17 dùng làm game ôn tập tích lũy.',
                    'source_preserved': True,
                }
                lesson['effective_kind'] = 'elective'
            if week <= 18:
                lessons.append(lesson)
            elif kind == 'assessment_feedback':
                lesson['sequence'] = None
                feedback.append(lesson)

        practice_scopes = []
        weeks = []
        for week in range(1, 19):
            weekly = [p for p in lessons if p['week'] == week]
            previous = [p for p in lessons if p['week'] < week and p['practice_eligible']]
            current = [p for p in weekly if p['practice_eligible']]
            scope = {
                'id': f'g{grade}-discover-w{week:02}', 'week': week,
                'prior_week_lesson_ids': [p['id'] for p in previous],
                'same_week_candidate_lesson_ids': [p['id'] for p in current],
                'prior_week_topics': list(dict.fromkeys(p['topic'] for p in previous)),
                'same_week_candidate_topics': list(dict.fromkeys(p['topic'] for p in current)),
                'same_week_order': 'Chưa có thời khóa biểu thực tế. Chỉ đưa vào câu hỏi sau khi GV xác nhận đã dạy phần tương ứng.',
            }
            practice_scopes.append(scope)
            for lesson in weekly:
                if lesson['effective_kind'] == 'elective':
                    lesson['elective_scope'] = {
                        'practice_scope_id': scope['id'],
                        'mode': 'cumulative_practice',
                        'same_week_earlier_lesson_ids': [p['id'] for p in current if p['row'] < lesson['row']],
                        'same_week_later_lesson_ids': [p['id'] for p in current if p['row'] > lesson['row']],
                        'same_week_order_verified': False,
                        'rule': 'Mặc định dùng tiết core từ các tuần trước. Cùng tuần: chỉ dùng tiết đứng trước trong PPCT và được GV xác nhận đã học; không lấy tiết đứng sau.',
                    }
            weeks.append({'week': week, 'count': len(weekly), 'periods': [p['period'] for p in weekly],
                          'kinds': dict(Counter(p['kind'] for p in weekly))})
        numbers = [p['period'] for p in lessons]
        duplicates = {p: [l['source']['period']['ref'] for l in lessons if l['period'] == p]
                      for p, count in Counter(numbers).items() if count > 1}
        missing = sorted(set(range(min(numbers), max(numbers) + 1)) - set(numbers))
        result['grades'].append({
            'grade': grade,
            'summary': {'actual_period_rows': len(lessons), 'source_first_period': min(numbers),
                        'source_last_period': max(numbers), 'kinds': dict(Counter(p['kind'] for p in lessons)),
                        'effective_kinds': dict(Counter(p['effective_kind'] for p in lessons)),
                        'duplicate_period_numbers': duplicates, 'missing_period_numbers': missing,
                        'content_without_period': no_period,
                        'unspecified_periods': [p['source']['period']['ref'] for p in lessons if p['kind'] == 'unspecified']},
            'weeks': weeks, 'lessons': lessons, 'practice_scopes': practice_scopes,
            'post_cutoff_feedback': feedback,
            'merged_ranges': [area for area in sheet['merges']
                              if range_boundaries(area)[1] <= lessons[-1]['row'] and range_boundaries(area)[0] <= 5],
        })
    validate(result)
    (OUT / 'ppct-hki.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({g['grade']: g['summary'] for g in result['grades']}, ensure_ascii=False, indent=2))


def validate(result):
    for grade in result['grades']:
        lessons = grade['lessons']
        by_id = {p['id']: p for p in lessons}
        assert len(lessons) == len(by_id)
        assert {p['week'] for p in lessons} == set(range(1, 19))
        assert len([p for p in lessons if p['kind'] == 'assessment' and p['week'] == 18]) == 2
        for p in lessons:
            assert p['source']['period']['anchor'] == f"D{p['row']}"
            assert p['period'] == p['source']['period']['cached']
            if p['effective_kind'] == 'elective':
                assert p['elective_scope']['mode'] == 'cumulative_practice'
        for scope in grade['practice_scopes']:
            for lesson_id in scope['prior_week_lesson_ids']:
                p = by_id[lesson_id]
                assert p['week'] < scope['week'] and p['kind'] == 'core'
            for lesson_id in scope['same_week_candidate_lesson_ids']:
                p = by_id[lesson_id]
                assert p['week'] == scope['week'] and p['kind'] == 'core'
    g10 = result['grades'][0]['lessons']
    assert [p['row'] for p in g10 if 11 <= p['row'] <= 13] == [11]
    assert next(p for p in g10 if p['row'] == 11)['source']['period']['range'] == 'D11:D13'
    assert [g['summary']['actual_period_rows'] for g in result['grades']] == [141, 137, 137]
    assert all(not p['practice_eligible'] for p in g10 if p['row'] in [117, 118])
    overrides = [p for p in result['grades'][2]['lessons'] if 'ownerOverride' in p]
    assert [p['period'] for p in overrides] == [121, 129]
    assert all(p['kind'] == 'unspecified' and p['effective_kind'] == 'elective'
               and p['source']['content']['raw'] is None for p in overrides)
    # Cross-check every exported period against the saved XML, independently of merged-cell resolution.
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    ranges = {'G10': (7, 170), 'G11': (13, 174), 'G12': (19, 182)}
    with ZipFile(SOURCE) as archive:
        rels = {e.attrib['Id']: e.attrib['Target'] for e in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        for sheet in ET.fromstring(archive.read('xl/workbook.xml')).find('s:sheets', ns):
            name = sheet.attrib['name']
            if name not in ranges:
                continue
            target = rels[sheet.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
            path = target.lstrip('/') if target.startswith('/') else 'xl/' + target
            low, high = ranges[name]
            expected = set()
            for cell in ET.fromstring(archive.read(path)).findall('.//s:sheetData/s:row/s:c', ns):
                coordinate = cell.attrib['r']
                value = cell.find('s:v', ns)
                if re.fullmatch(r'D\d+', coordinate) and low <= int(coordinate[1:]) <= high and value is not None:
                    expected.add((int(coordinate[1:]), int(value.text)))
            exported = {(p['row'], p['period']) for g in result['grades'] if f"G{g['grade']}" == name for p in g['lessons']}
            assert exported == expected, (name, exported ^ expected)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == result['source']['sha256']
    result['verification'] = {
        'source_sha256_unchanged': True,
        'independent_saved_xml_period_rows': [141, 137, 137],
        'all_numbered_period_rows_exported_once': True,
        'row_formula_matches_cached_period': True,
        'elective_scope_contains_no_future_week': True,
        'merged_period_children_not_duplicated': True,
    }


if __name__ == '__main__':
    if '--refresh' in sys.argv or not (OUT / 'sources' / 'ppct-raw.json').exists():
        inspect_workbook()
    normalize()
